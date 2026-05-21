"""调拨管理路由"""
import csv
import io
import os
import json
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Response
from app.api.deps import get_db, get_current_user, require_role
from app.schemas.warehouse import TransferCreate

router = APIRouter(prefix="/api/v1/transfers", tags=["调拨管理"])

# Status 映射（数据库中文 ↔ API 英文）
STATUS_TO_API = {'待审核': 'pending', '已通过': 'approved', '已驳回': 'rejected'}
STATUS_TO_DB = {v: k for k, v in STATUS_TO_API.items()}

# 调拨导入列名
TRANSFER_IMPORT_COLUMNS = ["物品名称", "源仓库", "目标仓库", "数量", "原因"]


def _generate_doc_no(conn) -> str:
    """生成调拨单号 DB-YYYYMMDD-XXX"""
    from datetime import datetime
    today = datetime.now().strftime("%Y%m%d")
    count = conn.execute(
        "SELECT COUNT(*) FROM transfers WHERE document_no LIKE ?",
        (f"DB-{today}-%",),
    ).fetchone()[0]
    return f"DB-{today}-{count + 1:03d}"


@router.get("/")
def list_transfers(
    page: int = 1,
    page_size: int = 20,
    status: str = "",
    current_user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """调拨记录列表"""
    where = "WHERE 1=1"
    params = []
    if status:
        where += " AND t.status = ?"
        params.append(STATUS_TO_DB.get(status, status))

    total = conn.execute(f"SELECT COUNT(*) FROM transfers t {where}", params).fetchone()[0]
    offset = (page - 1) * page_size

    rows = conn.execute(
        f"""SELECT t.*, i.name AS item_name,
           fw.name AS from_warehouse_name, tw.name AS to_warehouse_name,
           cu.display_name AS created_by_name, au.display_name AS approved_by_name
           FROM transfers t
           JOIN items i ON t.item_id = i.id
           JOIN warehouses fw ON t.from_warehouse_id = fw.id
           JOIN warehouses tw ON t.to_warehouse_id = tw.id
           JOIN users cu ON t.created_by = cu.id
           LEFT JOIN users au ON t.approved_by = au.id
           {where} ORDER BY t.id DESC LIMIT ? OFFSET ?""",
        params + [page_size, offset],
    ).fetchall()

    transfers = []
    for r in rows:
        d = dict(r)
        d["status"] = STATUS_TO_API.get(d["status"], d["status"])
        transfers.append(d)

    return {
        "transfers": transfers,
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@router.post("/")
def create_transfer(
    body: TransferCreate,
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """创建调拨申请"""
    if body.from_warehouse_id == body.to_warehouse_id:
        raise HTTPException(status_code=400, detail="源仓库和目标仓库不能相同")

    item = conn.execute("SELECT id, name FROM items WHERE id = ?", (body.item_id,)).fetchone()
    if not item:
        raise HTTPException(status_code=400, detail="物品不存在")

    from_wh = conn.execute("SELECT id FROM warehouses WHERE id = ?", (body.from_warehouse_id,)).fetchone()
    to_wh = conn.execute("SELECT id FROM warehouses WHERE id = ?", (body.to_warehouse_id,)).fetchone()
    if not from_wh or not to_wh:
        raise HTTPException(status_code=400, detail="仓库不存在")

    stock = conn.execute(
        "SELECT quantity FROM warehouse_stocks WHERE item_id = ? AND warehouse_id = ?",
        (body.item_id, body.from_warehouse_id),
    ).fetchone()
    available = stock["quantity"] if stock else 0
    if available < body.quantity:
        raise HTTPException(status_code=400, detail=f"源仓库库存不足（可用：{available}，需要：{body.quantity}）")

    doc_no = _generate_doc_no(conn)
    conn.execute(
        """INSERT INTO transfers (item_id, from_warehouse_id, to_warehouse_id, quantity, reason, document_no, created_by)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (body.item_id, body.from_warehouse_id, body.to_warehouse_id, body.quantity, body.reason, doc_no, current_user["id"]),
    )
    conn.commit()
    return {"message": "调拨申请已创建", "document_no": doc_no}


@router.post("/batch")
def create_batch_transfer(
    from_warehouse_id: int = Form(...),
    to_warehouse_id: int = Form(...),
    reason: str = Form(default=""),
    items: str = Form(default=""),
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """批量创建调拨申请（同一单据号，多物品）"""
    import json

    if from_warehouse_id == to_warehouse_id:
        raise HTTPException(status_code=400, detail="源仓库和目标仓库不能相同")

    try:
        items_list = json.loads(items)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="物品数据格式错误")

    if not items_list:
        raise HTTPException(status_code=400, detail="请至少添加一个物品")

    from_wh = conn.execute("SELECT id FROM warehouses WHERE id = ?", (from_warehouse_id,)).fetchone()
    to_wh = conn.execute("SELECT id FROM warehouses WHERE id = ?", (to_warehouse_id,)).fetchone()
    if not from_wh or not to_wh:
        raise HTTPException(status_code=400, detail="仓库不存在")

    doc_no = _generate_doc_no(conn)
    created = 0

    try:
        conn.execute("BEGIN IMMEDIATE")
        for item_data in items_list:
            item_id = int(item_data.get("item_id", 0))
            qty = int(item_data.get("quantity", 1))
            if qty < 1:
                continue

            item = conn.execute("SELECT id FROM items WHERE id = ?", (item_id,)).fetchone()
            if not item:
                continue

            stock = conn.execute(
                "SELECT quantity FROM warehouse_stocks WHERE item_id = ? AND warehouse_id = ?",
                (item_id, from_warehouse_id),
            ).fetchone()
            available = stock["quantity"] if stock else 0
            if available < qty:
                raise HTTPException(status_code=400,
                    detail=f"物品ID={item_id} 库存不足（可用：{available}，需要：{qty}）")

            conn.execute(
                """INSERT INTO transfers (item_id, from_warehouse_id, to_warehouse_id, quantity, reason, document_no, created_by)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (item_id, from_warehouse_id, to_warehouse_id, qty, reason, doc_no, current_user["id"]),
            )
            created += 1
        conn.commit()
    except HTTPException:
        conn.rollback()
        raise
    except Exception:
        conn.rollback()
        raise HTTPException(status_code=500, detail="批量调拨创建失败，已回滚")

    return {"message": f"调拨单已创建，共{created}件物品", "document_no": doc_no, "count": created}


@router.put("/{transfer_id}/approve")
def approve_transfer(
    transfer_id: int,
    current_user: dict = Depends(require_role("admin")),
    conn=Depends(get_db),
):
    """审核通过 → 执行库存调拨"""
    t = conn.execute("SELECT * FROM transfers WHERE id = ?", (transfer_id,)).fetchone()
    if not t:
        raise HTTPException(status_code=404, detail="调拨记录不存在")
    if t["status"] != "待审核":
        raise HTTPException(status_code=400, detail="该调拨记录已处理")

    # 再次校验库存
    stock = conn.execute(
        "SELECT quantity FROM warehouse_stocks WHERE item_id = ? AND warehouse_id = ?",
        (t["item_id"], t["from_warehouse_id"]),
    ).fetchone()
    available = stock["quantity"] if stock else 0
    if available < t["quantity"]:
        raise HTTPException(status_code=400, detail=f"源仓库库存不足（当前可用：{available}）")

    try:
        conn.execute("BEGIN IMMEDIATE")
        # 扣减源仓库
        conn.execute(
            "UPDATE warehouse_stocks SET quantity = quantity - ? WHERE item_id = ? AND warehouse_id = ?",
            (t["quantity"], t["item_id"], t["from_warehouse_id"]),
        )
        # 增加目标仓库
        existing = conn.execute(
            "SELECT id FROM warehouse_stocks WHERE item_id = ? AND warehouse_id = ?",
            (t["item_id"], t["to_warehouse_id"]),
        ).fetchone()
        if existing:
            conn.execute(
                "UPDATE warehouse_stocks SET quantity = quantity + ? WHERE item_id = ? AND warehouse_id = ?",
                (t["quantity"], t["item_id"], t["to_warehouse_id"]),
            )
        else:
            conn.execute(
                "INSERT INTO warehouse_stocks (item_id, warehouse_id, quantity) VALUES (?, ?, ?)",
                (t["item_id"], t["to_warehouse_id"], t["quantity"]),
            )
        # 更新调拨状态
        conn.execute(
            "UPDATE transfers SET status = '已通过', approved_by = ?, updated_at = datetime('now','localtime') WHERE id = ?",
            (current_user["id"], transfer_id),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise HTTPException(status_code=500, detail="调拨执行失败，已回滚")

    # 更新物品 total_quantity
    _update_item_total(conn, t["item_id"])
    conn.commit()

    return {"message": "调拨已通过，库存已转移"}


@router.put("/{transfer_id}/reject")
def reject_transfer(
    transfer_id: int,
    current_user: dict = Depends(require_role("admin")),
    conn=Depends(get_db),
):
    """驳回调拨"""
    t = conn.execute("SELECT * FROM transfers WHERE id = ?", (transfer_id,)).fetchone()
    if not t:
        raise HTTPException(status_code=404, detail="调拨记录不存在")
    if t["status"] != "待审核":
        raise HTTPException(status_code=400, detail="该调拨记录已处理")

    conn.execute(
        "UPDATE transfers SET status = '已驳回', approved_by = ?, updated_at = datetime('now','localtime') WHERE id = ?",
        (current_user["id"], transfer_id),
    )
    conn.commit()
    return {"message": "调拨已驳回"}


def _update_item_total(conn, item_id: int):
    """更新 items.total_quantity = 所有仓库库存之和"""
    total = conn.execute(
        "SELECT COALESCE(SUM(quantity), 0) FROM warehouse_stocks WHERE item_id = ?", (item_id,)
    ).fetchone()[0]
    conn.execute(
        "UPDATE items SET total_quantity = ?, updated_at = datetime('now','localtime') WHERE id = ?",
        (total, item_id),
    )


# ── 调拨批量导入 ──

def _parse_transfer_file(file_content: bytes, filename: str) -> list[dict]:
    """解析上传的调拨导入文件"""
    rows = []
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".csv":
        content = io.StringIO(file_content.decode("utf-8-sig"))
        reader = csv.DictReader(content)
        for row in reader:
            rows.append(row)
    elif ext in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_content), read_only=True)
        ws = wb.active
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            row_dict = {}
            for i, header in enumerate(headers):
                row_dict[header] = str(row[i]) if row[i] is not None else ""
            rows.append(row_dict)
        wb.close()
    else:
        raise HTTPException(status_code=400, detail="不支持的文件格式")
    return rows


def _validate_transfer_row(conn, row: dict, index: int) -> dict:
    """验证单行调拨数据"""
    name = row.get("物品名称", "").strip()
    from_wh = row.get("源仓库", "").strip()
    to_wh = row.get("目标仓库", "").strip()
    qty_str = row.get("数量", "1").strip()
    reason = row.get("原因", "").strip()

    if not name:
        return {"index": index, "data": row, "status": "error", "message": "物品名称不能为空"}
    if not from_wh:
        return {"index": index, "data": row, "status": "error", "message": "源仓库不能为空"}
    if not to_wh:
        return {"index": index, "data": row, "status": "error", "message": "目标仓库不能为空"}
    if from_wh == to_wh:
        return {"index": index, "data": row, "status": "error", "message": "源仓库和目标仓库不能相同"}

    try:
        qty = int(qty_str) if qty_str else 1
        if qty < 1:
            raise ValueError
    except ValueError:
        return {"index": index, "data": row, "status": "error", "message": f"数量必须为大于等于1的整数（当前：{qty_str}）"}

    item = conn.execute("SELECT id, name FROM items WHERE name = ?", (name,)).fetchone()
    if not item:
        return {"index": index, "data": row, "status": "error", "message": f"物品「{name}」不存在"}

    fw = conn.execute("SELECT id, name FROM warehouses WHERE name = ?", (from_wh,)).fetchone()
    if not fw:
        return {"index": index, "data": row, "status": "error", "message": f"源仓库「{from_wh}」不存在"}

    tw = conn.execute("SELECT id, name FROM warehouses WHERE name = ?", (to_wh,)).fetchone()
    if not tw:
        return {"index": index, "data": row, "status": "error", "message": f"目标仓库「{to_wh}」不存在"}

    stock = conn.execute(
        "SELECT quantity FROM warehouse_stocks WHERE item_id = ? AND warehouse_id = ?",
        (item["id"], fw["id"]),
    ).fetchone()
    available = stock["quantity"] if stock else 0
    if available < qty:
        return {
            "index": index, "data": row, "status": "error",
            "message": f"源仓库库存不足（{from_wh} 中「{name}」库存：{available}，需要：{qty}）",
        }

    return {
        "index": index, "data": row, "status": "ok", "message": None,
        "resolved": {"item_id": item["id"], "item_name": item["name"],
                     "from_warehouse_id": fw["id"], "from_warehouse_name": fw["name"],
                     "to_warehouse_id": tw["id"], "to_warehouse_name": tw["name"],
                     "quantity": qty, "reason": reason},
    }


@router.post("/import/preview")
def transfer_import_preview(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """上传调拨文件并预览"""
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in (".csv", ".xlsx", ".xls"):
        raise HTTPException(status_code=400, detail="仅支持 CSV 和 Excel 文件")

    file_content = file.file.read()
    if len(file_content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件大小不能超过 5MB")

    try:
        rows = _parse_transfer_file(file_content, file.filename)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败：{str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="文件中没有数据行")
    if len(rows) > 200:
        raise HTTPException(status_code=400, detail="单次导入最多 200 行")

    result = [_validate_transfer_row(conn, row, i + 1) for i, row in enumerate(rows)]
    summary = {
        "total": len(result),
        "ok": sum(1 for r in result if r["status"] == "ok"),
        "error": sum(1 for r in result if r["status"] == "error"),
    }
    return {"rows": result, "summary": summary}


@router.post("/import/confirm")
def transfer_import_confirm(
    file: UploadFile = File(...),
    selections: str = Form(default=""),
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """确认导入调拨单（事务保护）"""
    try:
        selected = json.loads(selections)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="选择数据格式错误")
    if not selected:
        raise HTTPException(status_code=400, detail="没有选择任何行")

    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in (".csv", ".xlsx", ".xls"):
        raise HTTPException(status_code=400, detail="仅支持 CSV 和 Excel 文件")

    file_content = file.file.read()
    if len(file_content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件大小不能超过 5MB")

    try:
        rows = _parse_transfer_file(file_content, file.filename)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败：{str(e)}")

    validated = [_validate_transfer_row(conn, row, i + 1) for i, row in enumerate(rows)]
    errors = [r for r in validated if r["status"] == "error"]
    if errors:
        msgs = "; ".join(f"第{r['index']}行: {r['message']}" for r in errors)
        raise HTTPException(status_code=400, detail=f"数据验证失败：{msgs}")

    selection_set = {s["index"] for s in selected}
    created = 0

    try:
        conn.execute("BEGIN IMMEDIATE")
        for row_data in validated:
            if row_data["index"] not in selection_set:
                continue
            res = row_data["resolved"]
            doc_no = _generate_doc_no(conn)
            conn.execute(
                """INSERT INTO transfers (item_id, from_warehouse_id, to_warehouse_id, quantity, reason, document_no, created_by, status)
                   VALUES (?, ?, ?, ?, ?, ?, ?, '待审核')""",
                (res["item_id"], res["from_warehouse_id"], res["to_warehouse_id"],
                 res["quantity"], res["reason"], doc_no, current_user["id"]),
            )
            created += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise HTTPException(status_code=500, detail="导入失败，已回滚")

    skipped = len(validated) - created
    return {"message": f"成功创建 {created} 条调拨申请", "imported": created, "skipped": skipped}


@router.get("/import/template")
def transfer_import_template(
    current_user: dict = Depends(get_current_user),
):
    """下载调拨导入模板"""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(TRANSFER_IMPORT_COLUMNS)
    writer.writerow(["示例：投影仪", "默认仓库", "A区主仓库", "3", "调拨原因示例"])
    content = output.getvalue()
    return Response(
        content="﻿" + content,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=transfer_import_template.csv"},
    )
