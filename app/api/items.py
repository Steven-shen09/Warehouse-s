"""物品管理路由"""
from typing import Optional
import csv
import io
import os
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, Query, Response
from app.api.deps import get_db, get_current_user, require_role
from app.schemas.item import ItemCreate, ItemUpdate, ImportPreviewRow, ImportConfirmRow, ImportConfirmRequest, ImportResult
from app.services.inventory_service import get_available_quantity, update_item_status

router = APIRouter(prefix="/api/v1/items", tags=["物品管理"])


# ── 导入模板列名映射 ──
IMPORT_COLUMNS = ["物品名称", "分类", "描述", "存放仓库", "总数量", "单价(元)", "预警阈值"]


def parse_uploaded_file(file_content: bytes, filename: str) -> list[dict]:
    """解析上传的 CSV/Excel 文件，返回行数据列表（每行一个 dict）"""
    rows = []
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".csv":
        content = io.StringIO(file_content.decode("utf-8-sig"))
        reader = csv.DictReader(content)
        for row in reader:
            rows.append(row)
    elif ext in (".xlsx", ".xls"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_content), read_only=True)
        ws = wb.active
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None for v in row):
                continue
            row_dict = {}
            for i, header in enumerate(headers):
                row_dict[header] = str(row[i]) if row[i] is not None else ""
            rows.append(row_dict)
        wb.close()
    else:
        raise HTTPException(status_code=400, detail="不支持的文件格式，请上传 CSV 或 Excel 文件")

    return rows


def validate_and_check_duplicates(conn, rows: list[dict]) -> list[dict]:
    """逐行验证并检查同名重复，返回预览行列表"""
    result = []
    for i, row in enumerate(rows):
        name = row.get("物品名称", "").strip()
        index = i + 1

        # 验证名称
        if not name:
            result.append({
                "index": index,
                "data": row,
                "status": "error",
                "message": "物品名称不能为空",
                "duplicate_item": None,
            })
            continue
        if len(name) > 100:
            result.append({
                "index": index,
                "data": row,
                "status": "error",
                "message": f"物品名称不能超过100字符（当前{len(name)}字符）",
                "duplicate_item": None,
            })
            continue

        # 验证数量
        qty_str = row.get("总数量", "1").strip()
        try:
            qty = int(qty_str) if qty_str else 1
            if qty < 1:
                raise ValueError
        except ValueError:
            result.append({
                "index": index,
                "data": row,
                "status": "error",
                "message": f"总数量必须为大于等于1的整数（当前值：{qty_str}）",
                "duplicate_item": None,
            })
            continue

        # 验证单价
        value_str = row.get("单价(元)", "0").strip()
        try:
            value = float(value_str) if value_str else 0.0
            if value < 0:
                raise ValueError
        except ValueError:
            result.append({
                "index": index,
                "data": row,
                "status": "error",
                "message": f"单价必须为非负数字（当前值：{value_str}）",
                "duplicate_item": None,
            })
            continue

        # 验证预警阈值
        threshold_str = row.get("预警阈值", "2").strip()
        try:
            threshold = int(threshold_str) if threshold_str else 2
            if threshold < 0:
                raise ValueError
        except ValueError:
            result.append({
                "index": index,
                "data": row,
                "status": "error",
                "message": f"预警阈值必须为非负整数（当前值：{threshold_str}）",
                "duplicate_item": None,
            })
            continue

        # 检查同名重复
        existing = conn.execute(
            "SELECT id, name, total_quantity FROM items WHERE name = ?", (name,)
        ).fetchone()

        if existing:
            result.append({
                "index": index,
                "data": row,
                "status": "duplicate",
                "message": f"同名物品已存在（ID:{existing['id']}，现有库存:{existing['total_quantity']}），勾选后累加",
                "duplicate_item": {"id": existing["id"], "name": existing["name"], "total_quantity": existing["total_quantity"]},
            })
        else:
            result.append({
                "index": index,
                "data": row,
                "status": "ok",
                "message": None,
                "duplicate_item": None,
            })

    return result


@router.get("/")
def list_items(
    page: int = 1,
    page_size: int = 20,
    keyword: str = "",
    category: str = "",
    status: str = "",
    low_stock: int = 0,
    warehouse_id: Optional[int] = Query(default=None),
    current_user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """物品列表（支持搜索、分类筛选、分页）"""
    where = "WHERE 1=1"
    params = []
    if keyword:
        where += " AND (name LIKE ? OR description LIKE ?)"
        params.extend([f"%{keyword}%", f"%{keyword}%"])
    if category:
        where += " AND category = ?"
        params.append(category)
    if status:
        where += " AND status = ?"
        params.append(status)
    if warehouse_id is not None and warehouse_id:
        where += " AND id IN (SELECT item_id FROM warehouse_stocks WHERE warehouse_id = ? AND quantity > 0)"
        params.append(warehouse_id)

    if low_stock:
        items = conn.execute(
            f"SELECT * FROM items {where} ORDER BY id DESC", params
        ).fetchall()
        from app.services.inventory_service import get_warehouse_stocks
        result = []
        for item in items:
            item_dict = dict(item)
            avail = get_available_quantity(conn, item["id"])
            item_dict["available_quantity"] = avail
            item_dict["warehouse_stocks"] = get_warehouse_stocks(conn, item["id"])
            if avail <= item["low_stock_threshold"]:
                result.append(item_dict)
        total = len(result)
        offset = (page - 1) * page_size
        paged = result[offset:offset + page_size]
        return {"items": paged, "total": total, "page": page, "page_size": page_size}

    total = conn.execute(f"SELECT COUNT(*) FROM items {where}", params).fetchone()[0]
    offset = (page - 1) * page_size

    items = conn.execute(
        f"SELECT * FROM items {where} ORDER BY id DESC LIMIT ? OFFSET ?",
        params + [page_size, offset],
    ).fetchall()

    from app.services.inventory_service import get_warehouse_stocks
    result = []
    for item in items:
        item_dict = dict(item)
        item_dict["available_quantity"] = get_available_quantity(conn, item["id"])
        item_dict["warehouse_stocks"] = get_warehouse_stocks(conn, item["id"])
        result.append(item_dict)

    return {"items": result, "total": total, "page": page, "page_size": page_size}


@router.get("/categories")
def get_categories(
    current_user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """获取所有物品分类列表"""
    rows = conn.execute("SELECT DISTINCT category FROM items WHERE category != '' ORDER BY category").fetchall()
    return {"categories": [r["category"] for r in rows]}


@router.get("/export")
def export_items(
    fmt: str = Query("csv", alias="format"),
    keyword: str = "",
    category: str = "",
    status: str = "",
    warehouse_id: Optional[int] = Query(default=None),
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """导出物品数据（不分页，支持筛选）"""
    where = "WHERE 1=1"
    params = []
    if keyword:
        where += " AND (name LIKE ? OR description LIKE ?)"
        params.extend([f"%{keyword}%", f"%{keyword}%"])
    if category:
        where += " AND category = ?"
        params.append(category)
    if status:
        where += " AND status = ?"
        params.append(status)
    if warehouse_id is not None and warehouse_id:
        where += " AND id IN (SELECT item_id FROM warehouse_stocks WHERE warehouse_id = ? AND quantity > 0)"
        params.append(warehouse_id)

    items = conn.execute(
        f"SELECT * FROM items {where} ORDER BY id DESC", params
    ).fetchall()

    from app.services.inventory_service import get_warehouse_stocks

    result = []
    for item in items:
        item_dict = dict(item)
        item_dict["available_quantity"] = get_available_quantity(conn, item["id"])
        stocks = get_warehouse_stocks(conn, item["id"])
        item_dict["warehouse_stocks"] = stocks
        item_dict["warehouse_stocks_str"] = "; ".join(f"{s['warehouse_name']}×{s['quantity']}" for s in stocks) if stocks else "-"
        result.append(item_dict)

    headers_row = ["ID", "名称", "分类", "描述", "总库存", "可用库存", "状态", "单价", "仓库库存分布"]

    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "物品数据"
        ws.append(headers_row)
        for item in result:
            ws.append([
                item["id"], item["name"], item["category"], item["description"],
                item["total_quantity"], item["available_quantity"],
                item["status"], item["value"], item["warehouse_stocks_str"],
            ])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=items_export.xlsx"},
        )
    else:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers_row)
        for item in result:
            writer.writerow([
                item["id"], item["name"], item["category"], item["description"],
                item["total_quantity"], item["available_quantity"],
                item["status"], item["value"], item["warehouse_stocks_str"],
            ])
        content = output.getvalue()
        return Response(
            content="﻿" + content,
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=items_export.csv"},
        )


@router.get("/{item_id}")
def get_item(
    item_id: int,
    current_user: dict = Depends(get_current_user),
    conn=Depends(get_db),
):
    """获取物品详情（含实时库存）"""
    item = conn.execute("SELECT * FROM items WHERE id = ?", (item_id,)).fetchone()
    if not item:
        raise HTTPException(status_code=404, detail="物品不存在")
    item_dict = dict(item)
    item_dict["available_quantity"] = get_available_quantity(conn, item_id)
    return item_dict


@router.post("/")
def create_item(
    body: ItemCreate,
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """新增物品"""
    conn.execute(
        """INSERT INTO items (name, category, description, total_quantity, value, low_stock_threshold, image_url)
           VALUES (?, ?, ?, ?, ?, ?, ?)""",
        (body.name, body.category, body.description,
         body.total_quantity, body.value, body.low_stock_threshold, body.image_url),
    )
    item_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]

    # 如果指定了仓库，写入 warehouse_stocks
    if body.warehouse_id:
        wh = conn.execute("SELECT id FROM warehouses WHERE id = ?", (body.warehouse_id,)).fetchone()
        if wh:
            conn.execute(
                "INSERT INTO warehouse_stocks (item_id, warehouse_id, quantity) VALUES (?, ?, ?)",
                (item_id, body.warehouse_id, body.total_quantity),
            )
    else:
        # 默认分配到默认仓库
        default_wh = conn.execute("SELECT id FROM warehouses ORDER BY id LIMIT 1").fetchone()
        if default_wh:
            conn.execute(
                "INSERT INTO warehouse_stocks (item_id, warehouse_id, quantity) VALUES (?, ?, ?)",
                (item_id, default_wh["id"], body.total_quantity),
            )

    conn.commit()
    return {"message": "物品创建成功"}


@router.put("/{item_id}")
def update_item(
    item_id: int,
    body: ItemUpdate,
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """更新物品信息"""
    item = conn.execute("SELECT id FROM items WHERE id = ?", (item_id,)).fetchone()
    if not item:
        raise HTTPException(status_code=404, detail="物品不存在")

    updates = {}
    for field in ["name", "category", "description", "total_quantity", "value", "low_stock_threshold", "image_url"]:
        val = getattr(body, field)
        if val is not None:
            updates[field] = val

    if updates:
        set_clause = ", ".join(f"{k} = ?" for k in updates)
        values = list(updates.values()) + [item_id]
        conn.execute(
            f"UPDATE items SET {set_clause}, updated_at = datetime('now','localtime') WHERE id = ?", values
        )
        conn.commit()
        # 更新库存状态
        update_item_status(conn, item_id)
        conn.commit()

    return {"message": "物品信息更新成功"}


@router.delete("/{item_id}")
def delete_item(
    item_id: int,
    current_user: dict = Depends(require_role("admin")),
    conn=Depends(get_db),
):
    """删除物品（前提：无活跃租借记录）"""
    active = conn.execute(
        "SELECT COUNT(*) FROM records WHERE item_id = ? AND status IN ('借出中', '逾期', '待审核')",
        (item_id,),
    ).fetchone()[0]
    if active > 0:
        raise HTTPException(status_code=400, detail="该物品存在活跃租借记录，无法删除")

    conn.execute("DELETE FROM items WHERE id = ?", (item_id,))
    conn.commit()
    return {"message": "物品已删除"}


@router.put("/{item_id}/mark-damaged")
def mark_damaged(
    item_id: int,
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """标记物品为损坏"""
    item = conn.execute("SELECT id FROM items WHERE id = ?", (item_id,)).fetchone()
    if not item:
        raise HTTPException(status_code=404, detail="物品不存在")
    conn.execute("UPDATE items SET status = '损坏', updated_at = datetime('now','localtime') WHERE id = ?", (item_id,))
    conn.commit()
    return {"message": "物品已标记为损坏"}


@router.put("/{item_id}/mark-available")
def mark_available(
    item_id: int,
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """恢复物品为可用"""
    item = conn.execute("SELECT id FROM items WHERE id = ?", (item_id,)).fetchone()
    if not item:
        raise HTTPException(status_code=404, detail="物品不存在")
    conn.execute("UPDATE items SET status = '可用', updated_at = datetime('now','localtime') WHERE id = ?", (item_id,))
    conn.commit()
    update_item_status(conn, item_id)
    conn.commit()
    return {"message": "物品已恢复为可用"}


# ── 批量导入 ──

@router.post("/import/preview")
def import_preview(
    file: UploadFile = File(...),
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """上传文件并预览解析结果（不写入数据库）"""
    # 验证文件类型
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in (".csv", ".xlsx", ".xls"):
        raise HTTPException(status_code=400, detail="仅支持 CSV (.csv) 和 Excel (.xlsx) 文件")

    # 验证文件大小（最大 5MB）
    file_content = file.file.read()
    if len(file_content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件大小不能超过 5MB")

    # 解析文件
    try:
        rows = parse_uploaded_file(file_content, file.filename)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败：{str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="文件中没有数据行")

    if len(rows) > 500:
        raise HTTPException(status_code=400, detail="单次导入最多 500 行")

    # 验证并检查重复
    preview_rows = validate_and_check_duplicates(conn, rows)

    # 统计
    summary = {
        "total": len(preview_rows),
        "ok": sum(1 for r in preview_rows if r["status"] == "ok"),
        "duplicate": sum(1 for r in preview_rows if r["status"] == "duplicate"),
        "error": sum(1 for r in preview_rows if r["status"] == "error"),
    }

    return {"rows": preview_rows, "summary": summary}


@router.post("/import/confirm")
def import_confirm(
    file: UploadFile = File(...),
    selections: str = Form(default=""),
    current_user: dict = Depends(require_role("admin", "approver")),
    conn=Depends(get_db),
):
    """确认导入（事务保护，全部或全不）"""
    import json

    try:
        selected_rows = json.loads(selections)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="选择数据格式错误")

    if not selected_rows:
        raise HTTPException(status_code=400, detail="没有选择任何行")

    # 验证文件类型
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in (".csv", ".xlsx", ".xls"):
        raise HTTPException(status_code=400, detail="仅支持 CSV (.csv) 和 Excel (.xlsx) 文件")

    # 验证文件大小（最大 5MB）
    file_content = file.file.read()
    if len(file_content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件大小不能超过 5MB")

    # 重新解析文件
    try:
        rows = parse_uploaded_file(file_content, file.filename)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"文件解析失败：{str(e)}")

    if not rows:
        raise HTTPException(status_code=400, detail="文件中没有数据行")

    if len(rows) > 500:
        raise HTTPException(status_code=400, detail="单次导入最多 500 行")

    # 重新验证
    preview_rows = validate_and_check_duplicates(conn, rows)
    errors = [r for r in preview_rows if r["status"] == "error"]
    if errors:
        error_msgs = "; ".join(f"第{r['index']}行: {r['message']}" for r in errors)
        raise HTTPException(status_code=400, detail=f"数据验证失败：{error_msgs}")

    # 构建选择的行索引查找
    selection_map = {s["index"]: s for s in selected_rows}

    imported = 0
    updated = 0

    try:
        conn.execute("BEGIN IMMEDIATE")

        for row_data in preview_rows:
            idx = row_data["index"]
            if idx not in selection_map:
                continue  # 跳过未勾选的行

            sel = selection_map[idx]
            data = row_data["data"]
            name = data.get("物品名称", "").strip()
            category = data.get("分类", "").strip()
            description = data.get("描述", "").strip()
            wh_name = data.get("存放仓库", "").strip()
            total_quantity = int(data.get("总数量", "1").strip() or "1")
            value = float(data.get("单价(元)", "0").strip() or "0")
            low_stock_threshold = int(data.get("预警阈值", "2").strip() or "2")

            # 解析仓库
            wh_id = None
            if wh_name:
                wh = conn.execute("SELECT id FROM warehouses WHERE name = ?", (wh_name,)).fetchone()
                if wh:
                    wh_id = wh["id"]
            if not wh_id:
                default_wh = conn.execute("SELECT id FROM warehouses ORDER BY id LIMIT 1").fetchone()
                if default_wh:
                    wh_id = default_wh["id"]

            if sel.get("action") == "add_to_existing":
                # 校验：该行必须是服务端确认的重复行
                if row_data["status"] != "duplicate" or not row_data.get("duplicate_item"):
                    raise HTTPException(status_code=400, detail=f"第{idx}行不是重复物品，无法累加")
                if sel.get("item_id") != row_data["duplicate_item"]["id"]:
                    raise HTTPException(status_code=400, detail=f"第{idx}行目标物品ID不匹配")
                # 验证目标物品存在
                target = conn.execute("SELECT id FROM items WHERE id = ?", (sel["item_id"],)).fetchone()
                if not target:
                    raise HTTPException(status_code=400, detail=f"目标物品(ID:{sel['item_id']})不存在")
                # 累加到已有物品
                conn.execute(
                    "UPDATE items SET total_quantity = total_quantity + ?, updated_at = datetime('now','localtime') WHERE id = ?",
                    (total_quantity, sel["item_id"]),
                )
                if wh_id:
                    existing_stock = conn.execute("SELECT id FROM warehouse_stocks WHERE item_id = ? AND warehouse_id = ?", (sel["item_id"], wh_id)).fetchone()
                    if existing_stock:
                        conn.execute("UPDATE warehouse_stocks SET quantity = quantity + ? WHERE item_id = ? AND warehouse_id = ?", (total_quantity, sel["item_id"], wh_id))
                    else:
                        conn.execute("INSERT INTO warehouse_stocks (item_id, warehouse_id, quantity) VALUES (?, ?, ?)", (sel["item_id"], wh_id, total_quantity))
                updated += 1
            else:
                # 创建新物品
                conn.execute(
                    """INSERT INTO items (name, category, description, total_quantity, value, low_stock_threshold)
                       VALUES (?, ?, ?, ?, ?, ?)""",
                    (name, category, description, total_quantity, value, low_stock_threshold),
                )
                new_item_id = conn.execute("SELECT last_insert_rowid()").fetchone()[0]
                if wh_id:
                    conn.execute("INSERT INTO warehouse_stocks (item_id, warehouse_id, quantity) VALUES (?, ?, ?)", (new_item_id, wh_id, total_quantity))
                imported += 1

        conn.commit()
    except HTTPException:
        conn.rollback()
        raise
    except Exception:
        conn.rollback()
        raise HTTPException(status_code=500, detail="导入失败，已回滚所有更改")

    skipped = len(preview_rows) - imported - updated

    return ImportResult(
        message=f"成功导入 {imported} 件物品，累加更新 {updated} 件物品",
        imported=imported,
        updated=updated,
        skipped=skipped,
    )


@router.get("/import/template")
def download_template(
    fmt: str = Query("csv", alias="format"),
    current_user: dict = Depends(get_current_user),
):
    """下载导入模板文件"""
    if fmt == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "物品导入模板"
        ws.append(IMPORT_COLUMNS)
        # 添加示例行
        ws.append(["示例：笔记本电脑", "电子设备", "ThinkPad X1", "默认仓库", "5", "4500", "2"])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=import_template.xlsx"},
        )
    else:
        # 默认 CSV
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(IMPORT_COLUMNS)
        content = output.getvalue()
        # 添加 BOM 以便 Excel 正确识别 UTF-8
        return Response(
            content="﻿" + content,
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=import_template.csv"},
        )


